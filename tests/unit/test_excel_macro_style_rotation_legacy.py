from __future__ import annotations

import hashlib
from datetime import date

import pandas as pd
import pytest
from openpyxl import Workbook

from alphahome.common.task_system.task_factory import UnifiedTaskFactory
from alphahome.fetchers.tasks.macro.excel_macro_style_rotation_legacy import (
    EXPECTED_INDICATOR_CODES,
    ExcelMacroStyleRotationLegacyTask,
    WORKBOOK_ENV_VAR,
    load_macro_style_rotation_legacy_workbook,
)


class _MockDB:
    async def get_column_names(self, target):
        return []

    async def fetch(self, query, *args, **kwargs):
        return []

    async def table_exists(self, target):
        return False


def _write_fixture(path, period: str, *, blank_electricity: bool = False):
    workbook = Workbook()
    memo = workbook.active
    memo.title = "Memo"
    memo.append(["指标", "分类", "指标_EN", "频率", "方向"])
    monthly = set(EXPECTED_INDICATOR_CODES[:-9])
    for code in EXPECTED_INDICATOR_CODES:
        memo.append(
            [
                code,
                "RATE" if code not in monthly else "GROWTH",
                code,
                "日" if code not in monthly else "月",
                1,
            ]
        )

    macro = workbook.create_sheet("CLEAN_MACRO")
    for row in range(1, 4):
        macro.cell(row, 1, None)
    macro.append(["日期", *EXPECTED_INDICATOR_CODES[:-9]])
    macro_values = [1.0] * len(EXPECTED_INDICATOR_CODES[:-9])
    if blank_electricity:
        macro_values[0] = None
    macro.append([pd.Timestamp(period).to_pydatetime(), *macro_values])

    rate = workbook.create_sheet("CLEAN_RATE")
    rate.append(["日期", *EXPECTED_INDICATOR_CODES[-9:]])
    rate.append([pd.Timestamp(period).to_pydatetime(), *([2.0] * 9)])
    workbook.save(path)


def _fixture_hash(path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def test_contract_contains_exactly_40_unique_indicators():
    assert len(EXPECTED_INDICATOR_CODES) == 40
    assert len(set(EXPECTED_INDICATOR_CODES)) == 40
    assert "tot_retail_sales_yoy" in EXPECTED_INDICATOR_CODES
    assert "retail_sales_yoy" not in EXPECTED_INDICATOR_CODES


def test_load_workbook_builds_complete_panel_and_pit_proxy(tmp_path):
    path = tmp_path / "legacy.xlsx"
    _write_fixture(path, "2011-02-28")

    frame = load_macro_style_rotation_legacy_workbook(
        path,
        "2011-02-28",
        "2011-02-28",
        expected_workbook_sha256=_fixture_hash(path),
    )

    assert len(frame) == 40
    assert frame["indicator_code"].nunique() == 40
    assert frame["is_observed"].all()
    assert frame["release_date"].isna().all()
    assert set(frame["strategy_available_date_proxy"]) == {date(2011, 4, 1)}
    assert frame["source_workbook_sha256"].str.fullmatch(r"[0-9a-f]{64}").all()
    tsf = frame.loc[frame["indicator_code"] == "TSF_yoy"].iloc[0]
    assert tsf["semantic_status"] == "legacy_mislabeled_rmb_loan_component"
    assert tsf["source_indicator_ids"] == "M004323990"


def test_january_blank_is_explicitly_structural(tmp_path):
    path = tmp_path / "legacy_january.xlsx"
    _write_fixture(path, "2012-01-31", blank_electricity=True)

    frame = load_macro_style_rotation_legacy_workbook(
        path,
        "2012-01-31",
        "2012-01-31",
        expected_workbook_sha256=_fixture_hash(path),
    )
    row = frame.loc[frame["indicator_code"] == "electricity_yoy"].iloc[0]

    assert bool(row["is_observed"]) is False
    assert row["missing_reason"] == "structural_january"


def test_contract_drift_fails_closed(tmp_path):
    path = tmp_path / "bad.xlsx"
    _write_fixture(path, "2011-02-28")
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    workbook["Memo"].cell(2, 3, "unexpected_indicator")
    workbook.save(path)

    with pytest.raises(ValueError, match="指标合同漂移"):
        load_macro_style_rotation_legacy_workbook(
            path,
            "2011-02-28",
            "2011-02-28",
            expected_workbook_sha256=_fixture_hash(path),
        )


def test_workbook_hash_drift_fails_closed(tmp_path):
    path = tmp_path / "drifted.xlsx"
    _write_fixture(path, "2011-02-28")

    with pytest.raises(ValueError, match="工作簿哈希不匹配"):
        load_macro_style_rotation_legacy_workbook(
            path,
            "2011-02-28",
            "2011-02-28",
            expected_workbook_sha256="0" * 64,
        )


def test_task_accepts_portable_environment_path(monkeypatch, tmp_path):
    path = tmp_path / "legacy-copy.xlsx"
    monkeypatch.setenv(WORKBOOK_ENV_VAR, str(path))

    task = ExcelMacroStyleRotationLegacyTask(db_connection=_MockDB())

    assert task.workbook_path == path


def test_task_is_registered_for_smart_execution():
    assert "excel_macro_style_rotation_legacy" in UnifiedTaskFactory._task_registry

    task = ExcelMacroStyleRotationLegacyTask(db_connection=_MockDB())
    assert task.data_source == "excel"
    assert task.date_column == "period_end_date"
    assert task.supports_incremental_update() is True
