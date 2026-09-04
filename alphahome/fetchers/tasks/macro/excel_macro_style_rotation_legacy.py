#!/usr/bin/env python

"""旧宏观风格轮动工作簿的可追溯历史快照。

该任务只负责冻结 2011-02 至 2020-12 的旧 Wind/同花顺 Excel 缓存值，
不把统计期伪装成发布日期。下游若要复现旧策略，应使用
``strategy_available_date_proxy``，并把真实官方源放在更高优先级。
"""

from __future__ import annotations

import asyncio
import hashlib
import math
import os
from collections.abc import Iterable
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, ClassVar

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from ....common.constants import UpdateTypes
from ....common.task_system.task_decorator import task_register
from ...base.fetcher_task import FetcherTask

WORKBOOK_ENV_VAR = "ALPHAHOME_MACRO_STYLE_ROTATION_WORKBOOK"
DEFAULT_WORKBOOK_PATH = (
    Path(__file__).resolve().parents[4].parent / "macroStrategy" / "宏观指标与逻辑.xlsx"
)
EXPECTED_WORKBOOK_SHA256 = (
    "7526814b13c4e4ada4d09639cf0ca3071e94e6549bebff84e1e8e57e1d3912b1"
)
LEGACY_START = pd.Timestamp("2011-02-28")
LEGACY_END = pd.Timestamp("2020-12-31")
CALCULATION_VERSION = "macro_strategy_excel_cache_v1"
AVAILABILITY_METHOD = "legacy_signal_delay_2_months"


@dataclass(frozen=True)
class LegacyIndicatorContract:
    indicator_code: str
    category: str
    source_frequency: str
    assumed_direction: int
    source_sheet: str


_MONTHLY_CODES = (
    "electricity_yoy",
    "electricity_ytd_yoy",
    "industrial_value_added_yoy",
    "fixedasset_investment_yoy",
    "fixedasset_investment_ytd_yoy",
    "pmi_manufacturing",
    "pmi_manufacturing_neworder",
    "pmi_nonmanufacturing",
    "export_yoy",
    "trade_balance_yoy",
    "CPI",
    "core_CPI",
    "PPI",
    "CPI_PPI",
    "core_CPI_PPI",
    "tot_retail_sales_yoy",
    "consumer_confidence",
    "M1",
    "M2",
    "M2_M1",
    "TSF_newadded_MA12_yoy",
    "TSF_yoy",
    "long_loan_newadded_MA12_yoy",
    "loan_yoy",
    "newstarts_area_yoy",
    "completed_yoy",
    "sold_area_yoy",
    "local_gov_budget_MA12_yoy",
    "newstarts_area_ytd_yoy",
    "completed_area_ytd_yoy",
    "sold_area_ytd_yoy",
)

_RATE_CODES = (
    "CN_BOND_1Y",
    "CN_BOND_10Y",
    "TERM_SPREAD",
    "US_BOND_5Y",
    "US_BOND_10Y",
    "CNY_INDEX",
    "USD_INDEX",
    "DR007",
    "CREDIT_SPREAD",
)

EXPECTED_INDICATOR_CODES = (*_MONTHLY_CODES, *_RATE_CODES)

_STRUCTURAL_JANUARY_CODES = {
    "electricity_yoy",
    "industrial_value_added_yoy",
    "fixedasset_investment_yoy",
    "fixedasset_investment_ytd_yoy",
    "tot_retail_sales_yoy",
    "newstarts_area_yoy",
    "completed_yoy",
    "sold_area_yoy",
    "newstarts_area_ytd_yoy",
    "completed_area_ytd_yoy",
    "sold_area_ytd_yoy",
}

_SOURCE_INDICATOR_IDS = {
    "electricity_yoy": "S002825946|S000006865",
    "electricity_ytd_yoy": "S000006865",
    "industrial_value_added_yoy": "M001622302|M001622303",
    "fixedasset_investment_yoy": "M001620537",
    "fixedasset_investment_ytd_yoy": "M001620538",
    "pmi_manufacturing": "M002043802",
    "pmi_manufacturing_neworder": "M002043804",
    "pmi_nonmanufacturing": "M002811179",
    "export_yoy": "M002808932",
    "trade_balance_yoy": "M004339261",
    "CPI": "M002826730",
    "core_CPI": "M004037496",
    "PPI": "M002826865",
    "CPI_PPI": "M002826730|M002826865",
    "core_CPI_PPI": "M004037496|M002826865",
    "tot_retail_sales_yoy": "M001625520|M001625521",
    "consumer_confidence": "M002807951",
    "M1": "M001625224",
    "M2": "M001625222",
    "M2_M1": "M001625222|M001625224",
    "TSF_newadded_MA12_yoy": "M004891015",
    "TSF_yoy": "M004323990",
    "long_loan_newadded_MA12_yoy": "M006017237",
    "loan_yoy": "M002859229",
    "newstarts_area_yoy": "S000047978",
    "completed_yoy": "S000047982",
    "sold_area_yoy": "S000047990",
    "local_gov_budget_MA12_yoy": "M004284835",
    "newstarts_area_ytd_yoy": "S000047978",
    "completed_area_ytd_yoy": "S000047982",
    "sold_area_ytd_yoy": "S000047990",
    "CN_BOND_1Y": "L001618296",
    "CN_BOND_10Y": "L001619604",
    "TERM_SPREAD": "L001619604|L001618296",
    "US_BOND_5Y": "G002600772",
    "US_BOND_10Y": "G002600774",
    "CNY_INDEX": "M004115462",
    "USD_INDEX": "G002600885",
    "DR007": "L001619493|THS:DR007.IB",
    "CREDIT_SPREAD": "L006741026",
}


def _workbook_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _normalise_frequency(value: Any) -> str:
    text = str(value or "").strip().lower()
    if text in {"月", "m", "monthly"}:
        return "monthly"
    if text in {"日", "d", "daily"}:
        return "daily_month_end"
    raise ValueError(f"无法识别指标频率: {value!r}")


def _load_contracts(workbook: Any) -> list[LegacyIndicatorContract]:
    sheet = workbook["Memo"]
    contracts: list[LegacyIndicatorContract] = []
    for row in sheet.iter_rows(min_row=2, max_col=5, values_only=True):
        category, indicator_code, frequency, direction = row[1:5]
        if indicator_code is None:
            continue
        code = str(indicator_code).strip()
        contracts.append(
            LegacyIndicatorContract(
                indicator_code=code,
                category=str(category).strip(),
                source_frequency=_normalise_frequency(frequency),
                assumed_direction=int(direction),
                source_sheet=(
                    "CLEAN_MACRO" if code in _MONTHLY_CODES else "CLEAN_RATE"
                ),
            )
        )

    actual = [item.indicator_code for item in contracts]
    if len(actual) != 40 or len(set(actual)) != 40:
        raise ValueError(f"Memo 指标必须恰好为 40 个且不能重复，当前为 {len(actual)}")
    if set(actual) != set(EXPECTED_INDICATOR_CODES):
        missing = sorted(set(EXPECTED_INDICATOR_CODES) - set(actual))
        unexpected = sorted(set(actual) - set(EXPECTED_INDICATOR_CODES))
        raise ValueError(
            f"Memo 指标合同漂移: missing={missing}, unexpected={unexpected}"
        )
    return contracts


def _coerce_numeric(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    try:
        numeric = float(value)
    except (TypeError, ValueError):
        return None
    return numeric if math.isfinite(numeric) else None


def _iter_sheet_rows(
    value_workbook: Any,
    formula_workbook: Any,
    contracts: Iterable[LegacyIndicatorContract],
    start_date: pd.Timestamp,
    end_date: pd.Timestamp,
    workbook_hash: str,
    source_workbook: str,
) -> list[dict[str, Any]]:
    by_sheet: dict[str, list[LegacyIndicatorContract]] = {
        "CLEAN_MACRO": [],
        "CLEAN_RATE": [],
    }
    for contract in contracts:
        by_sheet[contract.source_sheet].append(contract)

    rows: list[dict[str, Any]] = []
    for sheet_name, sheet_contracts in by_sheet.items():
        header_row = 4 if sheet_name == "CLEAN_MACRO" else 1
        first_data_row = header_row + 1
        value_sheet = value_workbook[sheet_name]
        formula_sheet = formula_workbook[sheet_name]
        header_cells = next(
            value_sheet.iter_rows(
                min_row=header_row, max_row=header_row, values_only=False
            )
        )
        headers = {
            str(cell.value).strip(): cell.column
            for cell in header_cells
            if cell.value is not None
        }
        missing_headers = sorted(
            item.indicator_code
            for item in sheet_contracts
            if item.indicator_code not in headers
        )
        if missing_headers:
            raise ValueError(f"{sheet_name} 缺少指标列: {missing_headers}")

        value_rows = value_sheet.iter_rows(min_row=first_data_row, values_only=True)
        formula_rows = formula_sheet.iter_rows(min_row=first_data_row, values_only=True)
        for row_number, (value_row, formula_row) in enumerate(
            zip(value_rows, formula_rows), first_data_row
        ):
            raw_period = value_row[0]
            if not isinstance(raw_period, (date, datetime, pd.Timestamp)):
                continue
            period = pd.Timestamp(raw_period).normalize()
            if period < start_date or period > end_date:
                continue
            for contract in sheet_contracts:
                column_number = headers[contract.indicator_code]
                cached_value = value_row[column_number - 1]
                value = _coerce_numeric(cached_value)
                formula_value = formula_row[column_number - 1]
                formula = (
                    formula_value
                    if isinstance(formula_value, str) and formula_value.startswith("=")
                    else None
                )
                rows.append(
                    {
                        "period_end_date": period.date(),
                        "indicator_code": contract.indicator_code,
                        "value": value,
                        "is_observed": value is not None,
                        "missing_reason": (
                            "excel_formula_error"
                            if isinstance(cached_value, str)
                            and cached_value.startswith("#")
                            else None
                        ),
                        "category": contract.category,
                        "source_frequency": contract.source_frequency,
                        "assumed_direction": contract.assumed_direction,
                        "release_date": None,
                        "strategy_available_date_proxy": (
                            period + pd.offsets.MonthBegin(2)
                        ).date(),
                        "availability_method": AVAILABILITY_METHOD,
                        "pit_lag_months": 2,
                        "source_name": "legacy_wind_excel_cache",
                        "source_workbook": source_workbook,
                        "source_workbook_sha256": workbook_hash,
                        "source_sheet": sheet_name,
                        "source_cell": f"{get_column_letter(column_number)}{row_number}",
                        "source_indicator_ids": _SOURCE_INDICATOR_IDS.get(
                            contract.indicator_code
                        ),
                        "source_formula": formula,
                        "calculation_version": CALCULATION_VERSION,
                        "semantic_status": (
                            "legacy_mislabeled_rmb_loan_component"
                            if contract.indicator_code == "TSF_yoy"
                            else "legacy_contract"
                        ),
                    }
                )
    return rows


def _complete_panel_and_missing_reasons(
    rows: list[dict[str, Any]], start_date: pd.Timestamp, end_date: pd.Timestamp
) -> pd.DataFrame:
    frame = pd.DataFrame(rows)
    if frame.empty:
        raise ValueError("旧工作簿在请求区间内没有数据")

    expected_periods = set(pd.date_range(start_date, end_date, freq="ME").date)
    expected_keys = {
        (period, code)
        for period in expected_periods
        for code in EXPECTED_INDICATOR_CODES
    }
    actual_keys = set(zip(frame["period_end_date"], frame["indicator_code"]))
    if actual_keys != expected_keys:
        missing = sorted(expected_keys - actual_keys)[:10]
        unexpected = sorted(actual_keys - expected_keys)[:10]
        raise ValueError(
            "旧工作簿未形成完整的 40 指标月度面板: "
            f"missing={missing}, unexpected={unexpected}"
        )

    first_observed = (
        frame.loc[frame["is_observed"]]
        .groupby("indicator_code")["period_end_date"]
        .min()
        .to_dict()
    )
    for index, row in frame.loc[~frame["is_observed"]].iterrows():
        current_reason = frame.at[index, "missing_reason"]
        code = row["indicator_code"]
        period = pd.Timestamp(row["period_end_date"])
        if code in _STRUCTURAL_JANUARY_CODES and period.month == 1:
            reason = "structural_january"
        elif code in first_observed and row["period_end_date"] < first_observed[code]:
            reason = (
                "insufficient_formula_history"
                if current_reason == "excel_formula_error"
                else "pre_series_inception"
            )
        elif pd.notna(current_reason) and str(current_reason).strip():
            reason = str(current_reason)
        else:
            reason = "source_blank"
        frame.at[index, "missing_reason"] = reason

    return frame.sort_values(["period_end_date", "indicator_code"]).reset_index(
        drop=True
    )


def load_macro_style_rotation_legacy_workbook(
    workbook_path: Path | str,
    start_date: pd.Timestamp | str = LEGACY_START,
    end_date: pd.Timestamp | str = LEGACY_END,
    expected_workbook_sha256: str | None = EXPECTED_WORKBOOK_SHA256,
) -> pd.DataFrame:
    path = Path(workbook_path)
    if not path.exists():
        raise FileNotFoundError(f"旧宏观策略工作簿不存在: {path}")
    start_ts = max(pd.Timestamp(start_date).normalize(), LEGACY_START)
    end_ts = min(pd.Timestamp(end_date).normalize(), LEGACY_END)
    if start_ts > end_ts:
        raise ValueError(
            f"请求区间不在旧工作簿范围内: {start_ts.date()}..{end_ts.date()}"
        )

    workbook_hash = _workbook_sha256(path)
    if (
        expected_workbook_sha256 is not None
        and workbook_hash.lower() != expected_workbook_sha256.lower()
    ):
        raise ValueError(
            "旧宏观策略工作簿哈希不匹配，拒绝覆盖冻结历史数据: "
            f"expected={expected_workbook_sha256.lower()}, actual={workbook_hash.lower()}"
        )
    value_workbook = load_workbook(
        path, read_only=True, data_only=True, keep_links=False
    )
    formula_workbook = load_workbook(
        path, read_only=True, data_only=False, keep_links=False
    )
    try:
        contracts = _load_contracts(value_workbook)
        rows = _iter_sheet_rows(
            value_workbook,
            formula_workbook,
            contracts,
            start_ts,
            end_ts,
            workbook_hash,
            path.name,
        )
        return _complete_panel_and_missing_reasons(rows, start_ts, end_ts)
    finally:
        value_workbook.close()
        formula_workbook.close()


@task_register()
class ExcelMacroStyleRotationLegacyTask(FetcherTask):
    domain = "macro"
    name = "excel_macro_style_rotation_legacy"
    description = "旧宏观风格轮动 40 指标历史快照（2011-02 至 2020-12）"
    table_name = "macro_style_rotation_legacy"
    data_source = "excel"
    primary_keys: ClassVar[list[str]] = ["period_end_date", "indicator_code"]
    date_column = "period_end_date"
    default_start_date = "20110228"
    update_type = UpdateTypes.SMART
    single_batch = True
    default_concurrent_limit = 1
    default_max_retries = 1
    default_save_batch_size = 5000

    schema_def: ClassVar[dict[str, dict[str, str]]] = {
        "period_end_date": {"type": "DATE", "constraints": "NOT NULL"},
        "indicator_code": {"type": "VARCHAR(64)", "constraints": "NOT NULL"},
        "value": {"type": "NUMERIC(28,10)"},
        "is_observed": {"type": "BOOLEAN", "constraints": "NOT NULL"},
        "missing_reason": {"type": "VARCHAR(40)"},
        "category": {"type": "VARCHAR(32)", "constraints": "NOT NULL"},
        "source_frequency": {
            "type": "VARCHAR(24)",
            "constraints": "NOT NULL",
        },
        "assumed_direction": {"type": "SMALLINT", "constraints": "NOT NULL"},
        "release_date": {
            "type": "DATE",
            "comment": "旧工作簿没有可信发布日期，必须为空",
        },
        "strategy_available_date_proxy": {
            "type": "DATE",
            "constraints": "NOT NULL",
            "comment": "旧回测口径：统计月后第二个月月初可用于交易",
        },
        "availability_method": {
            "type": "VARCHAR(64)",
            "constraints": "NOT NULL",
        },
        "pit_lag_months": {"type": "SMALLINT", "constraints": "NOT NULL"},
        "source_name": {"type": "VARCHAR(64)", "constraints": "NOT NULL"},
        "source_workbook": {"type": "TEXT", "constraints": "NOT NULL"},
        "source_workbook_sha256": {
            "type": "VARCHAR(64)",
            "constraints": "NOT NULL",
        },
        "source_sheet": {"type": "VARCHAR(32)", "constraints": "NOT NULL"},
        "source_cell": {"type": "VARCHAR(24)", "constraints": "NOT NULL"},
        "source_indicator_ids": {"type": "TEXT"},
        "source_formula": {"type": "TEXT"},
        "calculation_version": {
            "type": "VARCHAR(64)",
            "constraints": "NOT NULL",
        },
        "semantic_status": {
            "type": "VARCHAR(64)",
            "constraints": "NOT NULL",
        },
    }

    indexes: ClassVar[list[dict[str, Any]]] = [
        {
            "name": "idx_excel_macro_style_rotation_legacy_pk",
            "columns": "period_end_date, indicator_code",
            "unique": True,
        },
        {
            "name": "idx_excel_macro_style_rotation_legacy_indicator",
            "columns": "indicator_code, period_end_date",
        },
        {
            "name": "idx_excel_macro_style_rotation_legacy_available",
            "columns": "strategy_available_date_proxy",
        },
    ]

    validations: ClassVar[list[Any]] = [
        (lambda df: df["period_end_date"].notna(), "统计期不能为空"),
        (lambda df: df["indicator_code"].isin(EXPECTED_INDICATOR_CODES), "未知指标"),
        (
            lambda df: df["value"].isna()
            | df["value"].map(lambda value: math.isfinite(float(value))),
            "指标值必须为空或有限数",
        ),
        (lambda df: df["release_date"].isna(), "旧工作簿不得伪造发布日期"),
        (
            lambda df: pd.to_datetime(df["strategy_available_date_proxy"])
            > pd.to_datetime(df["period_end_date"]),
            "策略可用日代理必须晚于统计期",
        ),
        (
            lambda df: df["source_workbook_sha256"].str.fullmatch(
                r"[0-9a-f]{64}", na=False
            ),
            "工作簿哈希格式错误",
        ),
    ]

    def _apply_config(self, task_config: dict) -> None:
        super()._apply_config(task_config)
        configured_path = task_config.get("excel_file_path") or os.getenv(
            WORKBOOK_ENV_VAR
        )
        self.workbook_path = Path(configured_path or DEFAULT_WORKBOOK_PATH)

    def supports_incremental_update(self) -> bool:
        return True

    def get_incremental_skip_reason(self) -> str:
        return ""

    async def _determine_date_range(self) -> dict[str, str]:
        if self.update_type == UpdateTypes.MANUAL:
            if not self.start_date or not self.end_date:
                raise ValueError("Manual update requires start_date and end_date.")
            return {"start_date": self.start_date, "end_date": self.end_date}
        return {
            "start_date": LEGACY_START.strftime("%Y%m%d"),
            "end_date": LEGACY_END.strftime("%Y%m%d"),
        }

    async def get_batch_list(self, **kwargs) -> list[dict[str, Any]]:
        return [
            {
                "workbook_path": str(self.workbook_path),
                "start_date": kwargs["start_date"],
                "end_date": kwargs["end_date"],
            }
        ]

    async def prepare_params(self, batch: dict[str, Any]) -> dict[str, Any]:
        return batch.copy()

    async def fetch_batch(
        self,
        params: dict[str, Any],
        stop_event: asyncio.Event | None = None,
    ) -> pd.DataFrame:
        if stop_event and stop_event.is_set():
            raise asyncio.CancelledError
        return await asyncio.to_thread(
            load_macro_style_rotation_legacy_workbook,
            params["workbook_path"],
            params["start_date"],
            params["end_date"],
        )


__all__ = [
    "AVAILABILITY_METHOD",
    "CALCULATION_VERSION",
    "DEFAULT_WORKBOOK_PATH",
    "EXPECTED_WORKBOOK_SHA256",
    "EXPECTED_INDICATOR_CODES",
    "ExcelMacroStyleRotationLegacyTask",
    "LegacyIndicatorContract",
    "WORKBOOK_ENV_VAR",
    "load_macro_style_rotation_legacy_workbook",
]
