#!/usr/bin/env python

"""DR007 history recovered from the versioned iFinD workbook cache.

The workbook is a local vendor export.  It is useful for historical research
when the live Wind/iFinD APIs are unavailable, but it is deliberately labelled
as cached evidence and is not presented as a live official CFETS feed.
"""

from __future__ import annotations

import asyncio
import hashlib
import math
import os
from datetime import date, datetime
from pathlib import Path
from typing import Any, ClassVar

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from ....common.constants import UpdateTypes
from ....common.task_system.task_decorator import task_register
from ...base.fetcher_task import FetcherTask


WORKBOOK_ENV_VAR = "ALPHAHOME_MACRO_DR007_WORKBOOK"
DEFAULT_WORKBOOK_PATH = (
    Path(__file__).resolve().parents[4].parent / "macroStrategy" / "宏观指标与逻辑.xlsx"
)
SOURCE_SHEET = "DateRate"
SOURCE_SERIES_NAME = "DR007"
SOURCE_SERIES_ID = "L001619493"
SOURCE_UNIT = "%"
CALCULATION_VERSION = "ifind_excel_cache_dr007_v1"
AVAILABILITY_METHOD = "same_day_market_close_proxy"


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def load_dr007_workbook_cache(
    workbook_path: Path | str,
    start_date: date | datetime | str | pd.Timestamp | None = None,
    end_date: date | datetime | str | pd.Timestamp | None = None,
    expected_workbook_sha256: str | None = None,
) -> pd.DataFrame:
    """Read the daily DR007 column with source-cell and file-hash evidence."""

    path = Path(workbook_path)
    if not path.exists():
        raise FileNotFoundError(f"DR007 来源工作簿不存在: {path}")

    workbook_hash = _sha256(path)
    if expected_workbook_sha256 and workbook_hash.lower() != str(
        expected_workbook_sha256
    ).lower():
        raise ValueError(
            "DR007 来源工作簿哈希不匹配: "
            f"expected={str(expected_workbook_sha256).lower()}, actual={workbook_hash}"
        )

    start = pd.Timestamp(start_date).normalize() if start_date is not None else None
    end = pd.Timestamp(end_date).normalize() if end_date is not None else None
    if start is not None and end is not None and start > end:
        raise ValueError("DR007 请求起始日期晚于结束日期")

    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        if SOURCE_SHEET not in workbook.sheetnames:
            raise ValueError(f"DR007 来源工作簿缺少工作表: {SOURCE_SHEET}")
        sheet = workbook[SOURCE_SHEET]
        header_rows = list(
            sheet.iter_rows(min_row=1, max_row=4, values_only=True)
        )
        if len(header_rows) != 4:
            raise ValueError("DR007 来源工作簿表头不完整")

        series_columns = [
            index
            for index, value in enumerate(header_rows[1])
            if str(value or "").strip().upper() == SOURCE_SERIES_NAME
        ]
        if len(series_columns) != 1:
            raise ValueError(
                f"DR007 来源列必须恰好一列，当前匹配数: {len(series_columns)}"
            )
        column_index = series_columns[0]
        unit = str(header_rows[0][column_index] or "").strip()
        series_id = str(header_rows[3][column_index] or "").strip()
        if unit != SOURCE_UNIT or series_id != SOURCE_SERIES_ID:
            raise ValueError(
                "DR007 来源列元数据不符: "
                f"unit={unit!r}, series_id={series_id!r}"
            )

        observed_at = datetime.fromtimestamp(path.stat().st_mtime)
        rows: list[dict[str, Any]] = []
        for row_number, row in enumerate(
            sheet.iter_rows(min_row=5, values_only=True), start=5
        ):
            raw_date = row[0] if row else None
            raw_value = row[column_index] if len(row) > column_index else None
            if not isinstance(raw_date, (date, datetime, pd.Timestamp)):
                continue
            trade_date = pd.Timestamp(raw_date).normalize()
            if start is not None and trade_date < start:
                continue
            if end is not None and trade_date > end:
                continue
            try:
                value = float(raw_value)
            except (TypeError, ValueError):
                continue
            if not math.isfinite(value) or value < 0:
                continue
            rows.append(
                {
                    "trade_date": trade_date.date(),
                    "dr007_pct": value,
                    "availability_date_proxy": trade_date.date(),
                    "availability_method": AVAILABILITY_METHOD,
                    "is_weekend": trade_date.weekday() >= 5,
                    "source_vendor": "ifind_excel_cache",
                    "source_series_name": SOURCE_SERIES_NAME,
                    "source_series_id": SOURCE_SERIES_ID,
                    "source_unit": SOURCE_UNIT,
                    "source_workbook": path.name,
                    "source_workbook_sha256": workbook_hash,
                    "source_workbook_mtime": observed_at,
                    "source_sheet": SOURCE_SHEET,
                    "source_cell": f"{get_column_letter(column_index + 1)}{row_number}",
                    "evidence_status": "historical_vendor_cache_no_live_api",
                    "calculation_version": CALCULATION_VERSION,
                }
            )
    finally:
        workbook.close()

    if not rows:
        raise ValueError("DR007 来源工作簿在请求区间内没有有效数据")
    return (
        pd.DataFrame(rows)
        .sort_values(["trade_date", "source_cell"], kind="mergesort")
        .drop_duplicates("trade_date", keep="last")
        .reset_index(drop=True)
    )


@task_register()
class ExcelMacroDR007HistoryTask(FetcherTask):
    """Load cached daily DR007 values while retaining explicit provenance."""

    domain = "macro"
    name = "excel_macro_dr007_history"
    description = "iFinD 工作簿缓存 DR007 历史序列（带文件哈希证据）"
    table_name = "macro_dr007_history"
    data_source = "excel"
    primary_keys: ClassVar[list[str]] = ["trade_date"]
    date_column = "trade_date"
    default_start_date = "20141215"
    update_type = UpdateTypes.SMART
    single_batch = True
    default_concurrent_limit = 1
    default_max_retries = 1

    schema_def: ClassVar[dict[str, dict[str, str]]] = {
        "trade_date": {"type": "DATE", "constraints": "NOT NULL"},
        "dr007_pct": {"type": "NUMERIC(12,6)", "constraints": "NOT NULL"},
        "availability_date_proxy": {"type": "DATE", "constraints": "NOT NULL"},
        "availability_method": {"type": "VARCHAR(64)", "constraints": "NOT NULL"},
        "is_weekend": {"type": "BOOLEAN", "constraints": "NOT NULL"},
        "source_vendor": {"type": "VARCHAR(64)", "constraints": "NOT NULL"},
        "source_series_name": {"type": "VARCHAR(32)", "constraints": "NOT NULL"},
        "source_series_id": {"type": "VARCHAR(32)", "constraints": "NOT NULL"},
        "source_unit": {"type": "VARCHAR(16)", "constraints": "NOT NULL"},
        "source_workbook": {"type": "TEXT", "constraints": "NOT NULL"},
        "source_workbook_sha256": {"type": "VARCHAR(64)", "constraints": "NOT NULL"},
        "source_workbook_mtime": {"type": "TIMESTAMP", "constraints": "NOT NULL"},
        "source_sheet": {"type": "VARCHAR(64)", "constraints": "NOT NULL"},
        "source_cell": {"type": "VARCHAR(24)", "constraints": "NOT NULL"},
        "evidence_status": {"type": "VARCHAR(64)", "constraints": "NOT NULL"},
        "calculation_version": {"type": "VARCHAR(64)", "constraints": "NOT NULL"},
    }

    indexes: ClassVar[list[dict[str, Any]]] = [
        {
            "name": "idx_excel_macro_dr007_history_pk",
            "columns": "trade_date",
            "unique": True,
        },
        {
            "name": "idx_excel_macro_dr007_history_hash",
            "columns": "source_workbook_sha256",
        },
    ]

    validations: ClassVar[list[Any]] = [
        (lambda df: df["trade_date"].notna(), "日期不能为空"),
        (lambda df: df["dr007_pct"].between(0, 100), "DR007 必须为百分数口径"),
        (
            lambda df: df["availability_date_proxy"].eq(df["trade_date"]),
            "日度市场利率可用日代理必须等于交易日期",
        ),
        (
            lambda df: df["source_workbook_sha256"].str.fullmatch(
                r"[0-9a-f]{64}", na=False
            ),
            "工作簿哈希格式错误",
        ),
    ]

    def _apply_config(self, task_config: dict) -> None:
        super()._apply_config(task_config)
        configured_path = task_config.get("excel_file_path") or os.getenv(
            WORKBOOK_ENV_VAR
        )
        self.workbook_path = Path(configured_path or DEFAULT_WORKBOOK_PATH)
        self.expected_workbook_sha256 = task_config.get(
            "expected_workbook_sha256"
        )

    async def get_batch_list(self, **kwargs) -> list[dict[str, Any]]:
        return [
            {
                "workbook_path": str(self.workbook_path),
                "start_date": kwargs["start_date"],
                "end_date": kwargs["end_date"],
                "expected_workbook_sha256": self.expected_workbook_sha256,
            }
        ]

    async def prepare_params(self, batch: dict[str, Any]) -> dict[str, Any]:
        return batch.copy()

    async def fetch_batch(
        self,
        params: dict[str, Any],
        stop_event: asyncio.Event | None = None,
    ) -> pd.DataFrame:
        if stop_event and stop_event.is_set():
            raise asyncio.CancelledError
        return await asyncio.to_thread(
            load_dr007_workbook_cache,
            params["workbook_path"],
            params["start_date"],
            params["end_date"],
            params.get("expected_workbook_sha256"),
        )


__all__ = [
    "AVAILABILITY_METHOD",
    "CALCULATION_VERSION",
    "DEFAULT_WORKBOOK_PATH",
    "ExcelMacroDR007HistoryTask",
    "SOURCE_SERIES_ID",
    "WORKBOOK_ENV_VAR",
    "load_dr007_workbook_cache",
]
